"""
Timeline Analysis spreadsheet (Excel).

The same content as the Timeline Analysis HTML report (one source of
truth: reports.queries._query_timeline), as a workbook for sorting,
filtering and annotation:

  timeline : every timestamp of every revision of each document, in
             chronological order, with a SUMMARY row per document and a
             section per prior version
  summary  : one row per document: verdict, timestamps, flagged, the
             revisions examined, and timestamps assumed to be UTC

Rows 1-3 of each sheet carry the title, the provenance stamp (case,
generation time, version and build, source database; read by the case
index) and the UTC normalisation note.
"""

import os
import sqlite3
from datetime import datetime

from pdf_forensics import __version__, __tool_name__, BUILD_DATE
from pdf_forensics.reports.queries import _query_timeline

UTC_NOTE = ("All timestamps are normalised to UTC for comparison. Values "
            "with no timezone are assumed to be UTC (Timezone in source: "
            "none (assumed UTC)); an ordering problem involving one of them "
            "may reflect an unknown local offset of up to 14 hours.")


def generate_timeline_xlsx(db_path, output_path, session=None,
                           case_number=None):
    """Write the workbook. Returns the number of documents listed."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    conn = sqlite3.connect(db_path)
    try:
        columns, data = _query_timeline(conn)
    finally:
        conn.close()

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F3864")
    header_align = Alignment(horizontal="center", vertical="center",
                             wrap_text=True)
    cell_font = Font(size=10)
    bold = Font(size=10, bold=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")
    fills = {
        "FLAGGED": PatternFill("solid", fgColor="FDE2E1"),
        "INCONSISTENT": PatternFill("solid", fgColor="F8CBAD"),
        "CONSISTENT": PatternFill("solid", fgColor="E2EFDA"),
        "SUMMARY": PatternFill("solid", fgColor="EDEDED"),
    }
    stamp = (f"Case: {case_number or '-'}  |  Generated: "
             f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}  |  "
             f"v{__version__} (build {BUILD_DATE})  |  "
             f"Source DB: {os.path.basename(db_path)}")

    def head(ws, title, labels, widths):
        ws.cell(row=1, column=1,
                value=f"{__tool_name__} \u2014 {title}").font = \
            Font(bold=True, size=13)
        ws.cell(row=2, column=1, value=stamp).font = Font(size=9, italic=True)
        ws.cell(row=3, column=1, value=UTC_NOTE).font = \
            Font(size=9, italic=True, color="7F1D1D")
        for col, (lbl, w) in enumerate(zip(labels, widths), 1):
            c = ws.cell(row=5, column=col, value=lbl)
            c.font, c.fill, c.alignment, c.border = \
                header_font, header_fill, header_align, border
            ws.column_dimensions[get_column_letter(col)].width = w
        ws.freeze_panes = "B6"

    wb = Workbook()
    ws = wb.active
    ws.title = "timeline"
    names = [c["name"] for c in columns]
    labels = [c["label"] for c in columns]
    widths = {"Filename": 42, "#": 6, "Status": 15, "UTC": 20,
              "Revision": 34, "Timezone": 20, "Raw value": 28,
              "Category": 13, "Source": 34, "Field": 22, "Issue": 70,
              "Object": 14, "PDF MD5": 34}
    head(ws, "Timeline Analysis", labels, [widths.get(n, 16) for n in names])
    r = 6
    for rec in data:
        is_summary = rec.get("Source") == "SUMMARY"
        status = rec.get("Status", "")
        for col, n in enumerate(names, 1):
            v = rec.get(n, "")
            if n == "#" and is_summary:
                v = None
            c = ws.cell(row=r, column=col, value=v)
            c.font = bold if is_summary else cell_font
            c.border = border
            c.alignment = wrap if n in ("Issue", "Source", "Revision",
                                        "Raw value") else top
            if is_summary:
                c.fill = fills.get(status, fills["SUMMARY"])
            elif status == "FLAGGED":
                c.fill = fills["FLAGGED"]
        r += 1
    if r > 6:
        ws.auto_filter.ref = f"A5:{get_column_letter(len(names))}{r - 1}"

    # Summary sheet: one row per document.
    ws2 = wb.create_sheet("summary")
    s_labels = ["Filename", "Kind", "Verdict", "Timestamps", "Flagged",
                "Assumed UTC", "Detail", "PDF MD5"]
    head(ws2, "Timeline Analysis \u2014 summary", s_labels,
         [42, 22, 16, 12, 10, 12, 90, 34])
    docs = []
    for rec in data:
        if rec.get("Source") == "SUMMARY":
            docs.append({"rec": rec, "rows": []})
        elif docs:
            docs[-1]["rows"].append(rec)
    r = 6
    for d in docs:
        rec, rows = d["rec"], d["rows"]
        kind = ("Prior version" if rec["Issue"].startswith("Prior version")
                else "Submitted document")
        vals = [rec["Filename"], kind, rec["Status"], len(rows),
                sum(1 for x in rows if x.get("Status") == "FLAGGED"),
                sum(1 for x in rows
                    if x.get("Timezone") == "none (assumed UTC)"),
                rec["Issue"], rec["PDF MD5"]]
        for col, v in enumerate(vals, 1):
            c = ws2.cell(row=r, column=col, value=v)
            c.font, c.border = cell_font, border
            c.alignment = wrap if col == 7 else top
            if col == 3:
                c.fill = fills.get(rec["Status"], fills["SUMMARY"])
                c.font = bold
        r += 1
    if r > 6:
        ws2.auto_filter.ref = f"A5:H{r - 1}"

    wb.save(output_path)
    if session:
        session.log(f"Timeline Analysis spreadsheet: {len(docs)} document(s) "
                    f"-> {output_path}")
    return len(docs)
