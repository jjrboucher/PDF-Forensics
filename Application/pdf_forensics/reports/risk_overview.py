"""
Document Risk Overview spreadsheet.

Writes one row per analyzed document with its document-level risk score
and the per-category signal counts, in an Excel workbook built for
sorting and filtering (auto-filter + frozen header). This is the
spreadsheet counterpart of the Document Risk Dossier report, generated
from the database at report time.

Deliberately a NEW artifact rather than an edit of the triage sheet:
triage is a pre-analysis scan written before risk scores exist, and
rewriting a previously-saved spreadsheet after the fact would create an
artifact whose contents changed after its stated creation time. This
file is written once, at report time, from the current database state,
and carries its own generation stamp.
"""

import os
import sqlite3
from datetime import datetime

from pdf_forensics import __version__, __tool_name__, BUILD_DATE
from pdf_forensics.reports.queries import get_report_by_id


def generate_risk_overview_xlsx(db_path, output_path, session=None,
                                case_number=None):
    """
    Build the Document Risk Overview workbook from the risk_dossier
    query (single source of truth for the per-category columns).

    Returns the number of document rows written.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    rpt = get_report_by_id("risk_dossier")
    if rpt is None:
        return 0

    conn = sqlite3.connect(db_path)
    try:
        cur = conn.cursor()
        cur.execute(rpt["query"])
        headers = [d[0] for d in cur.description]
        rows = cur.fetchall()
    finally:
        conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "document-risk"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F3864")
    header_align = Alignment(horizontal="center", vertical="center",
                             wrap_text=True)
    cell_font = Font(size=10)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Title / provenance rows above the table.
    ws.cell(row=1, column=1,
            value=f"{__tool_name__} \u2014 Document Risk Overview").font = \
        Font(bold=True, size=13)
    stamp = (f"Case: {case_number or '-'}  |  Generated: "
             f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}  |  "
             f"v{__version__} (build {BUILD_DATE})  |  "
             f"Source DB: {os.path.basename(db_path)}")
    ws.cell(row=2, column=1, value=stamp).font = Font(size=9, italic=True)

    header_row = 4
    for col, hdr in enumerate(headers, 1):
        c = ws.cell(row=header_row, column=col, value=hdr)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align
        c.border = border

    for ri, row in enumerate(rows, header_row + 1):
        for col, val in enumerate(row, 1):
            c = ws.cell(row=ri, column=col, value=val)
            c.font = cell_font
            c.border = border

    # Column widths.
    for col, hdr in enumerate(headers, 1):
        width = max(12, min(40, len(str(hdr)) + 4))
        if str(hdr).lower() == "filename":
            width = 36
        elif "md5" in str(hdr).lower():
            width = 34
        ws.column_dimensions[get_column_letter(col)].width = width

    # Auto-filter over the header + data, frozen header, freeze Filename.
    last_col = get_column_letter(len(headers))
    last_row = header_row + len(rows)
    ws.auto_filter.ref = f"A{header_row}:{last_col}{max(last_row, header_row)}"
    ws.freeze_panes = f"B{header_row + 1}"

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    if session:
        session.log(f"  Document Risk Overview written: "
                    f"{os.path.basename(output_path)} ({len(rows)} documents)")
    return len(rows)
