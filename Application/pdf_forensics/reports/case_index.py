"""
Case index: a single landing page at the root of a case folder.

One self-contained HTML file ({case}_index.html) linking to every output
the application writes for a case: forensic reports, the visual overlay
report, relationship graphs, spreadsheets, extracted images, prior
versions, logs and databases.

Design
------
* A fixed sidebar menu switches between sections shown in one content
  area (no frames, no server, no external libraries), so the page works
  from local disk, a USB drive or a zipped folder. Each section has its
  own address (#reports etc.), so the back button and direct links work.
  Reports open in a new tab: each is a standalone exhibit with its own
  header, footer and print layout.
* Outputs are grouped by the ANALYSIS DATABASE they were produced from.
  Reports can be generated in a later session than the analysis, so the
  database is read from the file itself (every report, visual report and
  graph records it), not inferred from the filename timestamp. The latest
  database is shown by default; a selector shows earlier ones. Each
  output states the build that produced it, and the overview warns when
  a database's outputs come from more than one build.
* Case-level material that is not tied to one database (extracted
  images, prior versions, triage spreadsheets, logs) is always shown.
* A manifest records the size and MD5 of each report, spreadsheet,
  graph, database and log at the time the index was built. Extracted
  images are counted per folder rather than hashed individually.
  Hashes are cached by size and modification time so unchanged files are
  not re-hashed on every refresh.

The index is a navigation aid, not evidence: it is regenerated after
every operation that writes to the case folder and points to the
artifacts, which remain the authoritative records.
"""

import os
import re
import json
import glob
import html
import hashlib
import sqlite3
from datetime import datetime
from urllib.parse import quote

from pdf_forensics import __version__, __tool_name__, BUILD_DATE
from pdf_forensics import branding as _branding
from pdf_forensics.reports.queries import (
    REPORTS, REPORT_GROUPS, get_report_by_id,
    report_composed_of, report_part_of, is_aggregate_report,
)

_esc = html.escape

TS_RE = re.compile(r"_(\d{8}_\d{6})\.[A-Za-z0-9]+$")
REPORT_RE = re.compile(r"_report-([A-Za-z0-9_]+)_\d{8}_\d{6}\.html$")
GRAPH_RE = re.compile(r"_graph-([A-Za-z0-9_-]+)_\d{8}_\d{6}\.html$")
BUILD_RE = re.compile(rb"build (\d{4}-\d{2}-\d{2}[a-z]?)")  # optional same-day letter, e.g. 2026-10-22b
DB_HTML_RE = re.compile(rb"Database: ([^<\r\n]+?\.db)")
DB_GRAPH_RE = re.compile(rb'const DB_NAME = "([^"]+)"')
LOG_BUILD_RE = re.compile(r"\] Build: (\d{4}-\d{2}-\d{2}[a-z]?)")
LOG_DB_RE = re.compile(r"([^\s\\/]+_forensics_\d{8}_\d{6}\.db)")

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".jp2", ".jpx",
              ".bmp", ".gif", ".jbig2", ".jb2", ".webp", ".pbm", ".pgm",
              ".ppm", ".pam"}
CACHE_NAME = ".pdf_forensics_index_cache.json"
UNLINKED = "__unlinked__"


# --------------------------------------------------------------------------
# Scanning helpers
# --------------------------------------------------------------------------
def _ts_of(name):
    m = TS_RE.search(name)
    return m.group(1) if m else None


def _ts_display(ts):
    if not ts:
        return ""
    try:
        return datetime.strptime(ts, "%Y%m%d_%H%M%S").strftime(
            "%Y-%m-%d %H:%M:%S")
    except ValueError:
        return ts


def _read_head(path, limit=262144):
    try:
        with open(path, "rb") as fh:
            return fh.read(limit)
    except Exception:
        return b""


def _html_db_and_build(path, graph=False):
    """Source database name and build recorded inside an HTML output."""
    head = _read_head(path)
    db = build = None
    m = (DB_GRAPH_RE if graph else DB_HTML_RE).search(head)
    if not m and graph:
        # DB_NAME may be written late in a large graph file.
        try:
            if os.path.getsize(path) <= 60 * 1024 * 1024:
                with open(path, "rb") as fh:
                    full = fh.read()
                m = DB_GRAPH_RE.search(full)
                b = BUILD_RE.search(full)
                if b:
                    build = b.group(1).decode("ascii", "ignore")
        except Exception:
            pass
    if m:
        db = html.unescape(m.group(1).decode("utf-8", "ignore")).strip()
        db = os.path.basename(db)
    if build is None:
        b = BUILD_RE.search(head)
        if b:
            build = b.group(1).decode("ascii", "ignore")
    return db, build


def _xlsx_db_and_build(path):
    """Source DB and build from the risk-overview stamp row (A2)."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True)
        ws = wb.active
        stamp = str(ws["A2"].value or "")
        wb.close()
    except Exception:
        return None, None
    db = None
    m = re.search(r"Source DB: (\S+\.db)", stamp)
    if m:
        db = m.group(1)
    b = re.search(r"build (\d{4}-\d{2}-\d{2}[a-z]?)", stamp)
    return db, (b.group(1) if b else None)


def _log_info(path):
    """(build, [database names mentioned]) from a session log."""
    build = None
    dbs = []
    try:
        with open(path, encoding="utf-8", errors="ignore") as fh:
            for i, line in enumerate(fh):
                if build is None:
                    m = LOG_BUILD_RE.search(line)
                    if m:
                        build = m.group(1)
                for d in LOG_DB_RE.findall(line):
                    if d not in dbs:
                        dbs.append(d)
                if i > 20000:
                    break
    except Exception:
        pass
    return build, dbs


class _HashCache:
    def __init__(self, root):
        self.path = os.path.join(root, CACHE_NAME)
        self.root = root
        try:
            with open(self.path, encoding="utf-8") as fh:
                self.data = json.load(fh)
        except Exception:
            self.data = {}
        self.used = {}

    def md5(self, fpath):
        rel = os.path.relpath(fpath, self.root)
        try:
            st = os.stat(fpath)
        except OSError:
            return None
        key = [st.st_size, int(st.st_mtime)]
        hit = self.data.get(rel)
        if hit and hit[:2] == key:
            self.used[rel] = hit
            return hit[2]
        h = hashlib.md5()
        try:
            with open(fpath, "rb") as fh:
                for chunk in iter(lambda: fh.read(1 << 20), b""):
                    h.update(chunk)
        except Exception:
            return None
        digest = h.hexdigest()
        self.used[rel] = key + [digest]
        return digest

    def save(self):
        try:
            with open(self.path, "w", encoding="utf-8") as fh:
                json.dump(self.used, fh)
        except Exception:
            pass


def _file_rec(path, root, cache, kind, **extra):
    st = os.stat(path)
    name = os.path.basename(path)
    rec = {
        "kind": kind,
        "path": path,
        "rel": os.path.relpath(path, root).replace(os.sep, "/"),
        "name": name,
        "size": st.st_size,
        "mtime": st.st_mtime,
        "ts": _ts_of(name),
        "md5": cache.md5(path),
    }
    rec.update(extra)
    return rec


def _db_stats(path):
    """Read-only summary of an analysis database, including how the
    submitted files were accounted for (where recorded)."""
    out = {"docs": None, "pages": None, "findings": None,
           "first": None, "last": None, "originals": None,
           "versions": None, "runs": None, "submitted": None}
    try:
        uri = "file:" + quote(os.path.abspath(path).replace(os.sep, "/")) \
              + "?mode=ro"
        conn = sqlite3.connect(uri, uri=True)
        try:
            cur = conn.cursor()
            row = cur.execute(
                "SELECT COUNT(*), COALESCE(SUM(total_pages),0), "
                "MIN(analyzed_at), MAX(analyzed_at) FROM analyses").fetchone()
            out["docs"], out["pages"], out["first"], out["last"] = row
            out["findings"] = cur.execute(
                "SELECT COUNT(*) FROM findings").fetchone()[0]
            acols = {r[1] for r in cur.execute("PRAGMA table_info(analyses)")}
            if "version_number" in acols:
                out["versions"] = cur.execute(
                    "SELECT COUNT(*) FROM analyses "
                    "WHERE version_number IS NOT NULL").fetchone()[0]
                out["originals"] = out["docs"] - out["versions"]
            has_sub = cur.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' "
                "AND name='submissions'").fetchone()
            if has_sub:
                runs = {}
                for ts, outcome, n in cur.execute(
                        "SELECT session_ts, outcome, COUNT(*) FROM submissions "
                        "WHERE is_version = 0 GROUP BY session_ts, outcome"):
                    runs.setdefault(ts or "", {})[outcome] = n
                vers = {(ts or ""): n for ts, n in cur.execute(
                    "SELECT session_ts, COUNT(*) FROM submissions "
                    "WHERE is_version = 1 AND outcome = 'analysed' "
                    "GROUP BY session_ts")}
                out["runs"] = [(ts, runs[ts], vers.get(ts, 0))
                               for ts in sorted(runs)]
                out["submitted"] = cur.execute(
                    "SELECT pdf_filename, outcome, duplicate_of, pdf_md5, "
                    "reason FROM submissions WHERE is_version = 0 "
                    "ORDER BY id").fetchall()
        finally:
            conn.close()
    except Exception:
        pass
    return out


def _run_line(counts, versions):
    sub = sum(counts.values())
    parts = [f"{counts.get('analysed', 0)} analysed",
             f"{counts.get('duplicate', 0)} identical "
             f"cop{'y' if counts.get('duplicate', 0) == 1 else 'ies'} of "
             f"files already analysed",
             f"{counts.get('failed', 0)} failed"]
    if counts.get("encrypted_unopened"):
        parts.append(f"{counts['encrypted_unopened']} encrypted and not "
                     f"opened")
    if counts.get("not_found"):
        parts.append(f"{counts['not_found']} not found")
    line = f"{sub} submitted: " + ", ".join(parts)
    if versions:
        line += (f"; plus {versions} prior version"
                 f"{'' if versions == 1 else 's'} extracted and analysed")
    return line


def _analyses_text(st):
    if st.get("docs") is None:
        return "unknown"
    if st.get("versions"):
        return (f"{st['docs']} ({st['originals']} submitted "
                f"document{'' if st['originals'] == 1 else 's'} + "
                f"{st['versions']} prior version"
                f"{'' if st['versions'] == 1 else 's'})")
    return str(st["docs"])


def _submitted_html(st):
    runs = st.get("runs")
    if not runs:
        return ('<span class="muted">Not recorded \u2014 this database was '
                'analysed before submitted files were recorded (build '
                '2026-10-08). Identical copies skipped or files that failed '
                'in that run are not counted here; see the session log.'
                '</span>')
    lines = []
    for ts, counts, versions in runs:
        when = _ts_display(ts) if ts else ""
        lines.append(_esc(_run_line(counts, versions))
                     + (f' <span class="muted">(run {_esc(when)})</span>'
                        if len(runs) > 1 and when else ""))
    html_ = "<br>".join(lines)
    subs = st.get("submitted") or []
    if subs:
        rows = []
        for name, outcome, dup_of, md5, reason in subs:
            label = {"analysed": "Analysed",
                     "duplicate": "Identical copy \u2014 not re-analysed",
                     "failed": "Failed",
                     "encrypted_unopened": "Encrypted \u2014 not opened "
                                           "(password required)",
                     "not_found": "Not found"}.get(outcome, outcome)
            detail = (f"identical to {dup_of}" if outcome == "duplicate"
                      else (reason or ""))
            rows.append(f'<tr><td>{_esc(name)}</td><td>{_esc(label)}</td>'
                        f'<td>{_esc(detail)}</td><td class="mono nw">'
                        f'{_esc(md5 or "")}</td></tr>')
        html_ += (f'<details class="subs"><summary>Show all {len(subs)} '
                  f'submitted files</summary><div class="scroll"><table '
                  f'class="list"><tr><th>File</th><th>Outcome</th>'
                  f'<th>Detail</th><th>MD5</th></tr>{"".join(rows)}'
                  f'</table></div></details>')
    return html_


def _scan_case(root):
    cache = _HashCache(root)
    data = {"dbs": {}, "reports": [], "visual": [], "graphs": [],
            "triage": [], "risk_overview": [], "logs": [], "other": [],
            "images": [], "versions": []}

    # Databases at the case root.
    for p in sorted(glob.glob(os.path.join(glob.escape(root), "*.db"))):
        rec = _file_rec(p, root, cache, "database")
        rec["stats"] = _db_stats(p)
        rec["build"] = None
        data["dbs"][rec["name"]] = rec

    # Logs (also give each database its analysis build).
    logs_dir = os.path.join(root, "logs")
    for p in sorted(glob.glob(os.path.join(glob.escape(logs_dir), "*.log"))):
        build, dbs = _log_info(p)
        rec = _file_rec(p, root, cache, "log", build=build, dbs=dbs)
        data["logs"].append(rec)
    for db in data["dbs"].values():
        for lg in data["logs"]:
            if lg["ts"] and lg["ts"] == db["ts"] and lg["build"]:
                db["build"] = lg["build"]

    # Reports folder.
    rdir = os.path.join(root, "reports")
    for p in sorted(glob.glob(os.path.join(glob.escape(rdir), "*"))):
        if not os.path.isfile(p):
            continue
        name = os.path.basename(p)
        low = name.lower()
        if low.endswith(".html") and "_index" in low:
            continue                       # stale index from older builds
        m = REPORT_RE.search(name)
        if m:
            db, build = _html_db_and_build(p)
            data["reports"].append(_file_rec(p, root, cache, "report",
                                             rid=m.group(1), db=db,
                                             build=build))
        elif ("visual-overlays" in low or "visual_overlays" in low) \
                and low.endswith(".html"):
            db, build = _html_db_and_build(p)
            data["visual"].append(_file_rec(p, root, cache, "visual",
                                            db=db, build=build))
        elif "risk-overview" in low and low.endswith(".xlsx"):
            db, build = _xlsx_db_and_build(p)
            data["risk_overview"].append(_file_rec(
                p, root, cache, "risk_overview", db=db, build=build))
        elif "_triage" in low and low.endswith(".xlsx"):
            data["triage"].append(_file_rec(p, root, cache, "triage"))
        else:
            data["other"].append(_file_rec(p, root, cache, "other"))

    # Graphs.
    gdir = os.path.join(root, "graph_reports")
    for p in sorted(glob.glob(os.path.join(glob.escape(gdir), "*.html"))):
        m = GRAPH_RE.search(os.path.basename(p))
        db, build = _html_db_and_build(p, graph=True)
        data["graphs"].append(_file_rec(p, root, cache, "graph",
                                        scope=(m.group(1) if m else ""),
                                        db=db, build=build))

    # Extracted images: counted per folder, not hashed individually.
    idir = os.path.join(root, "extracted_images")
    if os.path.isdir(idir):
        for entry in sorted(os.listdir(idir)):
            sub = os.path.join(idir, entry)
            if not os.path.isdir(sub):
                continue
            n = size = 0
            latest = 0
            for dp, _dn, fns in os.walk(sub):
                for fn in fns:
                    if os.path.splitext(fn)[1].lower() in IMAGE_EXTS:
                        fp = os.path.join(dp, fn)
                        try:
                            st = os.stat(fp)
                        except OSError:
                            continue
                        n += 1
                        size += st.st_size
                        latest = max(latest, st.st_mtime)
            data["images"].append({
                "name": entry,
                "rel": os.path.relpath(sub, root).replace(os.sep, "/") + "/",
                "count": n, "size": size, "mtime": latest})

    # Prior versions.
    vdir = os.path.join(root, "pdf_versions")
    if os.path.isdir(vdir):
        for entry in sorted(os.listdir(vdir)):
            sub = os.path.join(vdir, entry)
            if not os.path.isdir(sub):
                continue
            files = sorted(glob.glob(os.path.join(glob.escape(sub), "*.pdf")))
            data["versions"].append({
                "name": re.sub(r"_versions$", "", entry),
                "rel": os.path.relpath(sub, root).replace(os.sep, "/") + "/",
                "files": [_file_rec(f, root, cache, "version")
                          for f in files]})

    # Link outputs to databases; fall back to a matching timestamp.
    known = set(data["dbs"])
    by_ts = {d["ts"]: n for n, d in data["dbs"].items() if d["ts"]}
    for key in ("reports", "visual", "graphs", "risk_overview"):
        for rec in data[key]:
            db = rec.get("db")
            if db not in known:
                db = by_ts.get(rec.get("ts"), db)
            rec["db"] = db if db in known else UNLINKED

    cache.save()
    return data


# --------------------------------------------------------------------------
# HTML helpers
# --------------------------------------------------------------------------
def _href(rel):
    return quote(rel, safe="/")


def _fmt_size(n):
    if n is None:
        return ""
    for unit in ("bytes", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return (f"{n:.0f} {unit}" if unit == "bytes"
                    else f"{n:.1f} {unit}")
        n /= 1024.0


def _fmt_time(t):
    return datetime.fromtimestamp(t).strftime("%Y-%m-%d %H:%M:%S") if t else ""


def _db_key(name):
    return re.sub(r"[^A-Za-z0-9_-]", "_", name)


def _meta_line(rec):
    bits = []
    if rec.get("ts"):
        bits.append(f"Generated {_ts_display(rec['ts'])}")
    elif rec.get("mtime"):
        bits.append(f"Modified {_fmt_time(rec['mtime'])}")
    if rec.get("build"):
        bits.append(f"build {rec['build']}")
    bits.append(_fmt_size(rec.get("size")))
    if rec.get("md5"):
        bits.append(f'MD5 <span class="mono" title="{rec["md5"]}">'
                    f'{rec["md5"][:12]}&hellip;</span>')
    return " &middot; ".join(bits)


def _desc_html(desc):
    """First sentence of a description, with the rest behind 'More'."""
    desc = (desc or "").strip()
    m = re.match(r"(.+?[.;])\s+(.+)", desc, re.S)
    if not m or len(desc) < 140:
        return f'<div class="rc-desc">{_esc(desc)}</div>'
    return (f'<div class="rc-desc">{_esc(m.group(1))}'
            f'<details class="more"><summary>More</summary>'
            f'{_esc(m.group(2))}</details></div>')


# File types a browser cannot display. Browsers will not launch desktop
# applications from a web page; a link to these only offers a download
# (and Chrome/Edge "open" a COPY placed in the Downloads folder). The
# index therefore offers "Copy path" for them instead of a link, so the
# examiner opens the original file in place.
NOT_VIEWABLE = {".xlsx", ".xlsm", ".xls", ".csv", ".db", ".sqlite"}


def _viewable(rec):
    return os.path.splitext(rec["name"])[1].lower() not in NOT_VIEWABLE


def _copy_btn(rec, label="Copy path"):
    return (f'<button type="button" class="copy-path" '
            f'data-rel="{_href(rec["rel"])}" title="Copy the full path to '
            f'this file, then paste it into Win+R or the File Explorer '
            f'address bar and press Enter to open the original file.">'
            f'{label}</button>')


_COPY_HINT = ('<span class="copy-hint">Browsers cannot open this file type '
              'directly. Copy the path, then paste it into Win+R or the File '
              'Explorer address bar and press Enter.</span>')


def _open_link(rec, label="Open"):
    return (f'<a class="open-link" href="{_href(rec["rel"])}" '
            f'target="_blank" rel="noopener">{label} &#8599;</a>')


def _file_card(rec, title, desc, tag=None):
    tag_html = (f'<span class="type-tag {tag[1]}">{tag[0]}</span>'
                if tag else "")
    return (f'<div class="card" data-state="present">'
            f'<div class="rc-head">{tag_html}'
            f'<span class="rc-name">{_esc(title)}</span></div>'
            f'{_desc_html(desc)}'
            f'<div class="rc-file">{_esc(rec["name"])}</div>'
            f'<div class="rc-status">'
            + (_open_link(rec) if _viewable(rec)
               else _copy_btn(rec) + _COPY_HINT)
            + f'<span class="gen-when">{_meta_line(rec)}</span></div></div>')


def _newest(recs):
    """Newest first. The session timestamp in the filename is set when the
    file is generated and survives copying; modification times are reset
    when a case folder is copied, so they are only a fallback."""
    return sorted(recs, key=lambda r: (r.get("ts") or "", r["mtime"]),
                  reverse=True)


def _report_cards(db_reports):
    """Catalogue cards for one database: every report, generated or not."""
    by_rid = {}
    for r in db_reports:
        by_rid.setdefault(r["rid"], []).append(r)
    present = set(by_rid)
    grouped = {rid for _l, ids in REPORT_GROUPS for rid in ids}
    groups = list(REPORT_GROUPS)
    left = [r["id"] for r in REPORTS if r["id"] not in grouped]
    if left:
        groups.append(("Other Reports", left))

    def rel_line(rid):
        out = []
        if is_aggregate_report(rid):
            names = []
            for cid in report_composed_of(rid):
                d = get_report_by_id(cid)
                lbl = _esc(d["name"] if d else cid)
                names.append(lbl if cid in present
                              else f'<span class="rel-missing">{lbl} '
                                   f'(not generated)</span>')
            out.append('<div class="rel combines">Combines: '
                       + ", ".join(names) + "</div>")
        parents = report_part_of(rid)
        if parents:
            names = [_esc(get_report_by_id(p)["name"]) for p in parents
                     if get_report_by_id(p)]
            out.append('<div class="rel partof">Also summarized in: '
                       + ", ".join(names) + "</div>")
        return "".join(out)

    parts = [f'<div class="summary-line">{len(present)} of {len(REPORTS)} '
             f'reports generated from this database.</div>']
    for label, ids in groups:
        cards = []
        for rid in ids:
            d = get_report_by_id(rid)
            if not d:
                continue
            agg = is_aggregate_report(rid)
            tag = ('<span class="type-tag agg">Overview</span>' if agg
                   else '<span class="type-tag det">Detail</span>')
            recs = _newest(by_rid.get(rid, []))
            if recs:
                cur = recs[0]
                status = (f'{_open_link(cur, "Open report")}'
                          f'<span class="gen-when">{_meta_line(cur)}</span>')
                if len(recs) > 1:
                    older = "".join(
                        f'<li><a href="{_href(o["rel"])}" target="_blank" '
                        f'rel="noopener">{_esc(o["name"])}</a> '
                        f'<span class="gen-when">{_meta_line(o)}</span></li>'
                        for o in recs[1:])
                    status += (f'<details class="older"><summary>'
                               f'{len(recs) - 1} earlier copy/copies'
                               f'</summary><ul>{older}</ul></details>')
                state = "present"
            else:
                status = '<span class="not-generated">Not generated</span>'
                state = "absent"
            cards.append(
                f'<div class="card{" aggregate" if agg else ""}" '
                f'data-state="{state}"><div class="rc-head">{tag}'
                f'<span class="rc-name">{_esc(d["name"])}</span></div>'
                f'{_desc_html(d.get("description", ""))}'
                f'{rel_line(rid)}<div class="rc-status">{status}</div></div>')
        if cards:
            parts.append(f'<h3 class="group-title">{_esc(label)}</h3>'
                         f'<div class="grid">{"".join(cards)}</div>')
    return "".join(parts)


def _empty(msg):
    return f'<div class="empty">{_esc(msg)}</div>'


# --------------------------------------------------------------------------
# Page assembly
# --------------------------------------------------------------------------
def _banner(case_number):
    cfg = _branding.load_branding()
    logo = _branding.get_agency_logo_data_uri(cfg)
    pairs = _branding.case_info_pairs(case_number, cfg)
    cls = _branding.classification_text(cfg)
    cls_html = f'<div id="classification">{_esc(cls)}</div>' if cls else ""
    logo_html = (f'<img class="agency-logo" src="{logo}" alt="Agency logo">'
                 if logo else "")
    info = ""
    if pairs:
        info = '<div class="case-info">' + "".join(
            f'<div class="bi-row"><span class="bi-label">{_esc(str(k))}'
            f'</span><span class="bi-val">{_esc(str(v))}</span></div>'
            for k, v in pairs) + "</div>"
    return (f'{cls_html}<div id="header"><div class="header-text">'
            f'<h1>Case Index</h1><div class="subtitle">'
            f'{_esc(__tool_name__)} v{_esc(__version__)} &middot; build '
            f'{_esc(BUILD_DATE)}<br>Index generated '
            f'{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}</div>'
            f'<div class="description">Start here. Every output for this '
            f'case is listed in the sections on the left. Reports open in '
            f'a new tab. Outputs are grouped by the analysis database they '
            f'were produced from; use the selector to view earlier '
            f'analyses.</div></div>{info}{logo_html}</div>')


def generate_case_index(case_root, case_number=None, session=None):
    """Build (or rebuild) {case}_index.html at the case root."""
    root = os.path.abspath(case_root)
    data = _scan_case(root)
    case = case_number or "case"

    dbs = sorted(data["dbs"].values(),
                 key=lambda d: (d["ts"] or "", d["mtime"]), reverse=True)
    db_keys = [d["name"] for d in dbs]
    has_unlinked = any(r["db"] == UNLINKED for k in
                       ("reports", "visual", "graphs", "risk_overview")
                       for r in data[k])
    if has_unlinked:
        db_keys.append(UNLINKED)

    def per_db(key, inner):
        k = _db_key(key)
        return f'<div class="per-db" data-db="{k}">{inner}</div>'

    def outs(kind, db):
        return [r for r in data[kind] if r["db"] == db]

    # ---------------- selector ----------------
    if db_keys:
        opts = []
        for i, k in enumerate(db_keys):
            if k == UNLINKED:
                lbl = "Outputs not linked to a database"
            else:
                d = data["dbs"][k]
                lbl = (f"{_ts_display(d['ts']) or k}"
                       + (" (latest)" if i == 0 else ""))
            opts.append(f'<option value="{_db_key(k)}"'
                        f'{" selected" if i == 0 else ""}>{_esc(lbl)}'
                        f'</option>')
        selector = (f'<div class="sel"><label for="dbsel">Analysis '
                    f'database</label><select id="dbsel">{"".join(opts)}'
                    f'</select></div>')
    else:
        selector = ('<div class="sel"><div class="empty">No analysis '
                    'database found in this folder yet.</div></div>')

    # ---------------- sections ----------------
    sec = {}

    # Overview
    ov = []
    for i, k in enumerate(db_keys):
        if k == UNLINKED:
            ov.append(per_db(k, _empty(
                "These outputs record a source database that is not in this "
                "folder (or none at all). They are listed in the other "
                "sections under this selection.")))
            continue
        d = data["dbs"][k]
        st = d["stats"]
        builds = {r["build"] for kind in ("reports", "visual", "graphs",
                                           "risk_overview")
                  for r in outs(kind, k) if r.get("build")}
        if d.get("build"):
            builds.add(d["build"])
        warn = ""
        if len(builds) > 1:
            warn = (f'<div class="warn">Outputs for this database were '
                    f'produced by more than one build '
                    f'({", ".join(sorted(builds))}). Results can differ '
                    f'between builds; check the build shown on each report '
                    f'and regenerate older outputs if consistency is '
                    f'required.</div>')
        rows = [("Database", _esc(d["name"]) + " " + _copy_btn(d)),
                ("Analysis session", _esc(_ts_display(d["ts"]))),
                ("Analysis build", _esc(d.get("build") or "not recorded "
                                        "(log predates build stamps)")),
                ("Submitted files", _submitted_html(st)),
                ("Analyses in database", _esc(_analyses_text(st))),
                ("Pages", _esc(str(st["pages"]))),
                ("Findings", _esc(str(st["findings"]))),
                ("Size / MD5", f'{_fmt_size(d["size"])} &middot; <span '
                               f'class="mono">{_esc(d["md5"] or "")}</span>')]
        table = "".join(f'<tr><th>{a}</th><td>{b}</td></tr>' for a, b in rows)
        quick = []
        for rid in ("risk_dossier", "concealment_overview"):
            recs = _newest([r for r in outs("reports", k) if r["rid"] == rid])
            dname = get_report_by_id(rid)["name"]
            quick.append(_file_card(recs[0], dname,
                                    get_report_by_id(rid)["description"],
                                    ("Overview", "agg")) if recs else
                         f'<div class="card" data-state="absent"><div '
                         f'class="rc-head"><span class="rc-name">'
                         f'{_esc(dname)}</span></div><div class="rc-status">'
                         f'<span class="not-generated">Not generated'
                         f'</span></div></div>')
        vis = _newest(outs("visual", k))
        if vis:
            quick.append(_file_card(vis[0], "Visual Overlay Report",
                                    "Each flagged page with its overlays "
                                    "as toggleable layers.",
                                    ("Visual", "vis")))
        rov = _newest(outs("risk_overview", k))
        if rov:
            quick.append(_file_card(rov[0], "Document Risk Overview",
                                    "Spreadsheet: one row per document, "
                                    "sortable by risk.", ("Excel", "xls")))
        ov.append(per_db(k,
            f'{warn}<h3 class="group-title">Selected analysis</h3>'
            f'<table class="kv">{table}</table>'
            f'<h3 class="group-title">Start here</h3>'
            f'<div class="grid">{"".join(quick)}</div>'))
    sess_rows = []
    for d in dbs:
        st = d["stats"]
        n_out = sum(len(outs(kind, d["name"])) for kind in
                    ("reports", "visual", "graphs", "risk_overview"))
        sess_rows.append(
            f'<tr><td><a href="#overview" class="pick" '
            f'data-pick="{_db_key(d["name"])}">'
            f'{_esc(_ts_display(d["ts"]) or d["name"])}</a></td>'
            f'<td class="mono">{_esc(d["name"])}</td>'
            f'<td>{_esc(_analyses_text(st))}</td>'
            f'<td>{_esc(d.get("build") or "not recorded")}</td>'
            f'<td>{n_out}</td></tr>')
    ov.append(
        '<h3 class="group-title">All analyses in this case folder</h3>'
        + ('<div class="scroll"><table class="list"><tr><th>Session</th>'
           '<th>Database</th><th>Analyses</th><th>Build</th>'
           '<th>Outputs</th></tr>' + "".join(sess_rows) + "</table></div>"
           if sess_rows else _empty("No analysis databases found.")))
    sec["overview"] = ("Overview", "".join(ov))

    # Forensic reports
    sec["reports"] = ("Forensic Reports", "".join(
        per_db(k, _report_cards(outs("reports", k))) for k in db_keys)
        or _empty("No analysis database found yet."))

    # Visual & graphs
    vg = []
    for k in db_keys:
        cards = [_file_card(r, "Visual Overlay Report",
                            "Flagged pages with overlays as toggleable "
                            "layers.", ("Visual", "vis"))
                 for r in _newest(outs("visual", k))]
        cards += [_file_card(r, f"Relationship Graph \u2014 "
                             f"{r['scope'] or 'graph'}",
                             "Image reuse across pages and documents.",
                             ("Graph", "gr"))
                  for r in _newest(outs("graphs", k))]
        vg.append(per_db(k, f'<div class="grid">{"".join(cards)}</div>'
                         if cards else _empty("No visual report or graphs "
                                              "generated from this "
                                              "database yet.")))
    sec["visual"] = ("Visual & Graphs", "".join(vg)
                     or _empty("Nothing generated yet."))

    # Spreadsheets
    sp = []
    for k in db_keys:
        cards = [_file_card(r, "Document Risk Overview",
                            "One row per document with risk score and "
                            "signal counts; sortable and filterable.",
                            ("Excel", "xls"))
                 for r in _newest(outs("risk_overview", k))]
        sp.append(per_db(k, f'<div class="grid">{"".join(cards)}</div>'
                         if cards else _empty("No risk overview "
                                              "spreadsheet for this "
                                              "database yet.")))
    tri = [_file_card(r, "Triage Spreadsheet",
                      "Fast pre-analysis scan: metadata and counts per PDF, "
                      "including prior-version rows.", ("Excel", "xls"))
           for r in _newest(data["triage"])]
    sp.append('<h3 class="group-title">Triage (all sessions)</h3>'
              + (f'<div class="grid">{"".join(tri)}</div>' if tri
                 else _empty("No triage spreadsheets in this case folder.")))
    sec["sheets"] = ("Spreadsheets", "".join(sp))

    # Extracted material
    em = ['<h3 class="group-title">Extracted images</h3>']
    if data["images"]:
        em.append('<div class="note">One folder per analysed document. '
                  'Folder contents reflect the most recent extraction for '
                  'that document.</div><div class="scroll"><table class="list">'
                  '<tr><th>Document</th><th>Images</th><th>Size</th>'
                  '<th>Last written</th></tr>')
        for f in data["images"]:
            em.append(f'<tr><td><a href="{_href(f["rel"])}" target="_blank" '
                      f'rel="noopener">{_esc(f["name"])}</a></td>'
                      f'<td>{f["count"]}</td><td>{_fmt_size(f["size"])}</td>'
                      f'<td>{_fmt_time(f["mtime"])}</td></tr>')
        em.append("</table></div>")
    else:
        em.append(_empty("No extracted images in this case folder."))
    em.append('<h3 class="group-title">Prior versions</h3>')
    if data["versions"]:
        for v in data["versions"]:
            items = "".join(
                f'<li><a href="{_href(f["rel"])}" target="_blank" '
                f'rel="noopener">{_esc(f["name"])}</a> <span '
                f'class="gen-when">{_meta_line(f)}</span></li>'
                for f in v["files"])
            em.append(f'<div class="card"><div class="rc-head"><span '
                      f'class="rc-name">{_esc(v["name"])}</span></div>'
                      f'<div class="rc-desc">{len(v["files"])} prior '
                      f'version(s) recovered from incremental updates.'
                      f'</div><ul class="files">{items}</ul></div>')
    else:
        em.append(_empty("No prior versions were extracted for this case."))
    sec["material"] = ("Extracted Material", "".join(em))

    # Logs & provenance
    lp = ['<h3 class="group-title">Session logs</h3>']
    if data["logs"]:
        lp.append('<div class="scroll"><table class="list"><tr><th>Log</th>'
                  '<th>Session</th><th>Build</th><th>Size</th><th>MD5</th>'
                  '</tr>')
        for r in _newest(data["logs"]):
            lp.append(f'<tr><td><a href="{_href(r["rel"])}" target="_blank" '
                      f'rel="noopener">{_esc(r["name"])}</a> '
                      f'{_copy_btn(r, "Copy path")}</td>'
                      f'<td class="nw">{_esc(_ts_display(r["ts"]))}</td>'
                      f'<td class="nw">{_esc(r.get("build") or "not recorded")}</td>'
                      f'<td class="nw">{_fmt_size(r["size"])}</td>'
                      f'<td class="mono nw">{_esc(r["md5"] or "")}</td></tr>')
        lp.append("</table></div>")
    else:
        lp.append(_empty("No session logs found."))
    manifest = []
    for kind_key, label in (("dbs", "Database"), ("reports", "Report"),
                            ("visual", "Visual report"), ("graphs", "Graph"),
                            ("risk_overview", "Spreadsheet"),
                            ("triage", "Spreadsheet"), ("logs", "Log"),
                            ("other", "Other")):
        recs = (list(data[kind_key].values()) if kind_key == "dbs"
                else data[kind_key])
        for r in recs:
            manifest.append((label, r))
    for v in data["versions"]:
        for r in v["files"]:
            manifest.append(("Prior version", r))
    lp.append('<h3 class="group-title">Output manifest</h3><div class="note">'
              'Size and MD5 of each output at the time this index was built '
              '(not a substitute for hashing the files at the point of '
              'export or hand-over). Extracted images are counted per folder '
              'in Extracted Material rather than listed here.</div>')
    if manifest:
        lp.append('<div class="scroll"><table class="list"><tr><th>Type</th>'
                  '<th>File</th><th>Database</th><th>Build</th>'
                  '<th>Modified</th><th>Size</th><th>MD5</th></tr>')
        for label, r in sorted(manifest, key=lambda x: x[1]["rel"]):
            dbn = r.get("db")
            dbn = ("not linked" if dbn == UNLINKED else (dbn or ""))
            fcell = (f'<a href="{_href(r["rel"])}" target="_blank" '
                     f'rel="noopener">{_esc(r["rel"])}</a>'
                     if _viewable(r) else
                     f'{_esc(r["rel"])} {_copy_btn(r, "Copy path")}')
            lp.append(f'<tr><td>{label}</td><td>{fcell}'
                      f'</td><td class="mono">{_esc(dbn)}</td>'
                      f'<td class="nw">{_esc(r.get("build") or "")}</td>'
                      f'<td class="nw">{_fmt_time(r["mtime"])}</td>'
                      f'<td class="nw">{_fmt_size(r["size"])}</td>'
                      f'<td class="mono nw">{_esc(r["md5"] or "")}</td></tr>')
        lp.append("</table></div>")
    else:
        lp.append(_empty("No outputs yet."))
    sec["provenance"] = ("Logs & Provenance", "".join(lp))

    # ---------------- assemble ----------------
    nav = "".join(f'<a href="#{sid}" data-sec="{sid}">{_esc(title)}</a>'
                  for sid, (title, _b) in sec.items())
    body = "".join(f'<section id="sec-{sid}" data-sec="{sid}">'
                   f'<h2>{_esc(title)}</h2>{content}</section>'
                   for sid, (title, content) in sec.items())
    page = f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Case Index &mdash; {_esc(case)}</title>
<style>{_CSS}</style></head><body>
{_banner(case_number)}
<div class="app"><nav class="sidebar">{selector}{nav}</nav>
<main>{body}</main></div>
<div id="index-footer">Case index generated by {_esc(__tool_name__)}
 v{_esc(__version__)} (build {_esc(BUILD_DATE)}) on
 {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}. This index is a
 navigation aid and is rebuilt after every operation; the linked outputs
 are the authoritative records.</div>
<script>{_JS}</script></body></html>"""

    out = os.path.join(root, f"{case.replace(os.sep, '_')}_index.html")
    with open(out, "w", encoding="utf-8") as fh:
        fh.write(page)

    # Remove the per-folder index written by earlier builds (our own
    # navigation file, superseded by this one).
    old = os.path.join(root, "reports", f"{case}_index.html")
    if os.path.isfile(old):
        try:
            os.remove(old)
        except OSError:
            pass

    if session:
        try:
            session.log(f"Case index written: {os.path.basename(out)}")
        except Exception:
            pass
    return out


def refresh_case_index(session=None, case_root=None, case_number=None,
                       log=None):
    """
    Rebuild the case index after any operation that writes outputs.
    Never raises: a failed refresh is logged and ignored, so it can never
    interrupt the operation that called it. Returns the path or None.
    """
    root = case_root or (getattr(session, "db_dir", None) if session else None)
    cn = case_number or (getattr(session, "case_number", None)
                         if session else None)
    if not root or not os.path.isdir(root):
        return None
    try:
        return generate_case_index(root, cn, session)
    except Exception as e:
        if log:
            try:
                log(f"Case index not refreshed: {e}")
            except Exception:
                pass
        if session:
            try:
                session.log(f"Case index not refreshed: {e}", "warning")
            except Exception:
                pass
        return None


_CSS = """
* { box-sizing: border-box; }
html, body { margin: 0; }
body { font-family: 'Segoe UI', Arial, sans-serif; background: #141414;
  color: #e8e8e8; }
#classification { background: #7a0000; color: #fff; text-align: center;
  font-weight: 700; padding: 4px; letter-spacing: 1px; }
#header { display: flex; align-items: center; gap: 20px; padding: 16px 24px;
  background: #1c1c1c; border-bottom: 3px solid #ED1C24; }
#header h1 { margin: 0; font-size: 24px; color: #fff; }
.header-text { flex: 1; }
.subtitle { color: #bdbdbd; font-size: 12px; margin-top: 3px; }
.description { color: #cfcfcf; font-size: 13px; margin-top: 8px;
  max-width: 860px; }
.agency-logo { max-height: 70px; max-width: 200px; }
.case-info { font-size: 12px; color: #ddd; }
.bi-row { display: flex; gap: 8px; }
.bi-label { color: #9a9a9a; min-width: 90px; }
.bi-val { color: #fff; font-weight: 600; }
.app { display: flex; align-items: stretch; min-height: 70vh; }
.sidebar { flex: 0 0 230px; background: #1a1a1a; border-right: 1px solid #2c2c2c;
  padding: 14px 10px; position: sticky; top: 0; align-self: flex-start;
  max-height: 100vh; overflow-y: auto; }
.sidebar a { display: block; color: #ddd; text-decoration: none;
  padding: 8px 10px; border-radius: 6px; font-size: 14px; margin: 2px 0; }
.sidebar a:hover { background: #262626; }
.sidebar a.active { background: #ED1C24; color: #fff; font-weight: 600; }
.sel { margin-bottom: 14px; padding-bottom: 12px; border-bottom: 1px solid #2c2c2c; }
.sel label { display: block; font-size: 11px; color: #9a9a9a;
  text-transform: uppercase; letter-spacing: .5px; margin-bottom: 5px; }
.sel select { width: 100%; background: #262626; color: #fff;
  border: 1px solid #444; border-radius: 5px; padding: 6px; font: inherit;
  font-size: 13px; }
main { flex: 1; padding: 18px 26px; min-width: 0; }
main section { display: none; }
main section.shown { display: block; }
main h2 { margin: 0 0 12px; color: #fff; font-size: 20px; }
.per-db { display: none; }
.per-db.shown { display: block; }
.group-title { color: #fff; font-size: 15px; border-bottom: 1px solid #333;
  padding-bottom: 5px; margin: 20px 0 10px; }
.summary-line, .note { color: #bdbdbd; font-size: 13px; margin-bottom: 10px; }
.grid { display: grid; gap: 12px;
  grid-template-columns: repeat(auto-fill, minmax(320px, 1fr)); }
.card { background: #1e1e1e; border: 1px solid #333; border-radius: 8px;
  padding: 12px 14px; margin-bottom: 10px; }
.grid .card { margin-bottom: 0; }
.card.aggregate { border-color: #ED1C24; background: #241b1b; }
.card[data-state="absent"] { opacity: .6; }
.rc-head { display: flex; align-items: center; gap: 8px; margin-bottom: 5px; }
.rc-name { font-weight: 700; color: #fff; font-size: 14px; }
.rc-desc { color: #c7c7c7; font-size: 12.5px; line-height: 1.45; }
.rc-file { color: #8f8f8f; font-size: 11px; margin-top: 5px;
  word-break: break-all; }
.type-tag { font-size: 10px; font-weight: 700; padding: 1px 6px;
  border-radius: 4px; letter-spacing: .5px; }
.type-tag.agg { background: #ED1C24; color: #fff; }
.type-tag.det { background: #2a2a2a; color: #bbb; border: 1px solid #444; }
.type-tag.vis { background: #4A90D9; color: #fff; }
.type-tag.gr { background: #8E7CC3; color: #fff; }
.type-tag.xls { background: #1D9E75; color: #fff; }
.rel { font-size: 11.5px; margin-top: 7px; color: #9fb4c7; }
.rel.combines { color: #e0a0a0; }
.rel-missing { color: #777; }
.rc-status { margin-top: 9px; display: flex; flex-wrap: wrap;
  align-items: center; gap: 10px; }
.open-link { color: #fff; background: #ED1C24; text-decoration: none;
  font-size: 12px; font-weight: 600; padding: 4px 10px; border-radius: 5px; }
.open-link:hover { background: #c01820; }
.gen-when { color: #8f8f8f; font-size: 11px; }
.copy-path { font: inherit; font-size: 12px; font-weight: 600; cursor: pointer;
  color: #fff; background: #1D9E75; border: none; border-radius: 5px;
  padding: 4px 10px; }
.copy-path:hover { background: #17835f; }
td .copy-path, table.kv .copy-path { font-size: 11px; padding: 2px 8px;
  margin-left: 6px; }
.copy-path.done { background: #4A90D9; }
.copy-hint { flex-basis: 100%; color: #9a9a9a; font-size: 11px; }
.not-generated { color: #8f8f8f; font-size: 12px; font-style: italic; }
.older { width: 100%; font-size: 12px; color: #bdbdbd; }
.older summary { cursor: pointer; }
.older a, .files a, .list a { color: #9fc5ff; }
.files { margin: 8px 0 0; padding-left: 18px; font-size: 12.5px; }
.files li { margin: 3px 0; }
.muted { color: #9a9a9a; }
.subs { margin-top: 6px; }
.subs summary { cursor: pointer; color: #9fc5ff; font-size: 12.5px; }
.subs table.list { margin-top: 6px; }
.empty { color: #8f8f8f; font-style: italic; font-size: 13px;
  padding: 10px 0; }
.warn { background: #3a2a10; border-left: 3px solid #EF9F27; color: #f2c48d;
  padding: 9px 12px; border-radius: 5px; font-size: 13px; margin-bottom: 12px; }
.mono { font-family: Consolas, Monaco, monospace; font-size: 11.5px; }
.scroll { overflow-x: auto; }
td.nw, th.nw { white-space: nowrap; }
.more { display: inline; }
.more summary { display: inline; cursor: pointer; color: #9fc5ff;
  font-size: 11.5px; margin-left: 4px; }
.more[open] summary { display: block; margin: 4px 0 2px; }
table.kv { border-collapse: collapse; font-size: 13px; }
table.kv th { text-align: left; color: #9a9a9a; font-weight: 600;
  padding: 4px 18px 4px 0; vertical-align: top; }
table.kv td { padding: 4px 0; color: #fff; }
table.list { border-collapse: collapse; font-size: 12.5px; width: 100%; }
table.list th { text-align: left; background: #262626; color: #ddd;
  padding: 6px 8px; position: sticky; top: 0; }
table.list td { padding: 5px 8px; border-top: 1px solid #2a2a2a;
  vertical-align: top; }
#index-footer { padding: 12px 24px; color: #777; font-size: 11px;
  border-top: 1px solid #2a2a2a; }
@media (max-width: 800px) {
  .app { flex-direction: column; }
  .sidebar { position: static; max-height: none; flex: none; width: 100%; }
}
@media print {
  body { background: #fff; color: #000; }
  .sidebar { display: none; }
  #header { background: #fff; }
  #header h1, .rc-name, .group-title, main h2, table.kv td { color: #000; }
  .card { background: #fff; border-color: #999; }
  .rc-desc, .subtitle, .description, table.list td { color: #000; }
  table.list th { background: #eee; color: #000; }
}
"""

_JS = r"""
(function () {
  var sel = document.getElementById('dbsel');
  function showDb(key) {
    document.querySelectorAll('.per-db').forEach(function (el) {
      el.classList.toggle('shown', el.getAttribute('data-db') === key);
    });
  }
  function showSection() {
    var id = (location.hash || '#overview').slice(1);
    var found = false;
    document.querySelectorAll('main section').forEach(function (s) {
      var on = s.getAttribute('data-sec') === id;
      s.classList.toggle('shown', on);
      if (on) found = true;
    });
    if (!found) {
      var first = document.querySelector('main section');
      if (first) first.classList.add('shown');
      id = first ? first.getAttribute('data-sec') : '';
    }
    document.querySelectorAll('.sidebar a[data-sec]').forEach(function (a) {
      a.classList.toggle('active', a.getAttribute('data-sec') === id);
    });
  }
  if (sel) {
    sel.addEventListener('change', function () { showDb(sel.value); });
    showDb(sel.value);
  }
  document.addEventListener('click', function (ev) {
    var p = ev.target.closest ? ev.target.closest('[data-pick]') : null;
    if (p && sel) { sel.value = p.getAttribute('data-pick'); showDb(sel.value); }
  });
  function localPath(rel) {
    var u = new URL(rel, location.href);
    var path = decodeURIComponent(u.pathname);
    if (u.host) {                        // UNC share: \\server\share\...
      return '\\\\' + u.host + path.replace(/\//g, '\\');
    }
    if (/^\/[A-Za-z]:\//.test(path)) {   // Windows drive: C:\...
      return path.slice(1).replace(/\//g, '\\');
    }
    return path;                         // macOS / Linux
  }
  function flash(btn, text) {
    var old = btn.getAttribute('data-label') || btn.textContent;
    btn.setAttribute('data-label', old);
    btn.textContent = text; btn.classList.add('done');
    setTimeout(function () {
      btn.textContent = old; btn.classList.remove('done');
    }, 1600);
  }
  function fallbackCopy(text) {
    var ta = document.createElement('textarea');
    ta.value = text; ta.style.position = 'fixed'; ta.style.opacity = '0';
    document.body.appendChild(ta); ta.select();
    var ok = false;
    try { ok = document.execCommand('copy'); } catch (e) {}
    document.body.removeChild(ta);
    return ok;
  }
  document.addEventListener('click', function (ev) {
    var btn = ev.target.closest ? ev.target.closest('.copy-path') : null;
    if (!btn) return;
    var path = localPath(btn.getAttribute('data-rel'));
    function manual() { window.prompt('Copy this path:', path); }
    if (navigator.clipboard && navigator.clipboard.writeText) {
      navigator.clipboard.writeText(path).then(function () {
        flash(btn, 'Copied');
      }, function () {
        if (fallbackCopy(path)) flash(btn, 'Copied'); else manual();
      });
    } else if (fallbackCopy(path)) {
      flash(btn, 'Copied');
    } else {
      manual();
    }
  });
  window.addEventListener('hashchange', showSection);
  showSection();
})();
"""
